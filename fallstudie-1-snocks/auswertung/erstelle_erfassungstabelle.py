"""Erzeugt die Erfassungstabelle für die Interviews (Repertory Grid, Laddering, Analogien).

Aufruf:  python3 erstelle_erfassungstabelle.py [Zielpfad]
Standard-Ziel: ../daten/Erfassung_Interviews.xlsx

Achtung: Überschreibt eine vorhandene Datei nur mit --force, damit keine
eingetragenen Interviewdaten verloren gehen.
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HIER = Path(__file__).resolve().parent
ZIEL = HIER.parent / "daten" / "Erfassung_Interviews.xlsx"

ELEMENTE = ["E1 SNOCKS", "E2 FALKE", "E3 Happy Socks", "E4 erlich textil",
            "E5 Calvin Klein", "E6 Uniqlo", "E7 Lidl-Eigenmarke", "E8 Ideal"]

GRID_KATEGORIEN = [
    ("Qualität & Haltbarkeit", "Material, Verarbeitung, hält lange, keine Löcher"),
    ("Komfort & Passform", "Tragegefühl, weich, sitzt gut, rutscht nicht, kratzt nicht"),
    ("Preis & Preis-Leistung", "günstig vs. teuer, fairer Preis, Mengenrabatt"),
    ("Nachhaltigkeit & Fairness", "Bio-Baumwolle, Produktion, Umwelt, faire Bedingungen"),
    ("Design & Stil", "bunt vs. schlicht, modisch, Muster, Optik"),
    ("Markenimage & Status", "Logo, Prestige, cool, angesagt, Luxus"),
    ("Bekanntheit & Vertrauen", "kenne ich, verlässlich, seriös, Garantie"),
    ("Einkauf & Verfügbarkeit", "online vs. Laden, überall erhältlich, bequem bestellbar"),
    ("Modernität & Innovation", "jung, neu, Start-up, digital"),
    ("Tradition & Herkunft", "Traditionsmarke, deutsch, Handwerk"),
    ("Sonstiges", "passt in keine Kategorie"),
]

CODES = {
    "A – Attribut": ["Bio-Baumwolle", "Anti-Loch-Garantie", "Multipacks / fairer Preis",
                     "Passform / Komfort", "hochwertiges Material", "schlichtes Design",
                     "buntes Design", "Markenlogo", "nachhaltige Produktion", "online bestellbar",
                     "günstiger Preis"],
    "FK – funktionale Konsequenz": ["angenehmes Tragegefühl", "hält lange / keine Löcher",
                                    "sitzt gut / rutscht nicht", "kratzt nicht", "atmungsaktiv",
                                    "spart Geld", "spart Zeit beim Einkauf", "kein Risiko beim Kauf",
                                    "umweltschonend hergestellt", "passt zu allem", "sieht gut aus"],
    "PK – psychosoziale Konsequenz": ["fühle mich wohl in meiner Haut", "muss nicht darüber nachdenken",
                                      "gutes Gewissen", "sehe gepflegt aus", "fühle mich selbstbewusst",
                                      "smart eingekauft", "Ärger vermeiden", "gehöre dazu",
                                      "kann mich auf Wichtiges konzentrieren"],
    "W – Wert": ["Lebensqualität / Wohlbefinden", "Sicherheit", "Selbstbewusstsein / Selbstachtung",
                 "Verantwortung / Umweltschutz", "Freiheit / Unabhängigkeit", "Effizienz / Zeit für Wichtiges",
                 "Sparsamkeit", "Anerkennung", "Spaß / Lebensfreude", "Zugehörigkeit"],
}

ARIAL = "Arial"
F_TITEL = Font(name=ARIAL, size=14, bold=True)
F_KOPF = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
F_TEXT = Font(name=ARIAL, size=10)
F_BEISPIEL = Font(name=ARIAL, size=10, italic=True, color="808080")
FILL_KOPF = PatternFill("solid", fgColor="333333")
FILL_FORMEL_KOPF = PatternFill("solid", fgColor="5B6B7F")
FILL_EINGABE = PatternFill("solid", fgColor="FFF6D5")
FILL_BEISPIEL = PatternFill("solid", fgColor="EDEDED")
RAND = Border(*(Side(style="thin", color="D0D0D0"),) * 4)
UMBRUCH = Alignment(wrap_text=True, vertical="top")

ZEILEN = {"Teilnehmende": 20, "Grid": 160, "Laddering": 60, "Analogien": 20}


def kopf(ws, spalten, breiten, formel_ab=None):
    for i, (titel, breite) in enumerate(zip(spalten, breiten), start=1):
        c = ws.cell(row=1, column=i, value=titel)
        c.font = F_KOPF
        c.fill = FILL_FORMEL_KOPF if formel_ab and i >= formel_ab else FILL_KOPF
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = breite
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "B3"


def beispiel(ws, werte):
    for i, w in enumerate(werte, start=1):
        if w is None:
            continue
        c = ws.cell(row=2, column=i, value=w)
        c.font = F_BEISPIEL
        c.fill = FILL_BEISPIEL
        c.alignment = UMBRUCH


def eingabe(ws, n_spalten, n_zeilen, bis_spalte=None):
    bis = bis_spalte or n_spalten
    for r in range(3, 3 + n_zeilen):
        for col in range(1, n_spalten + 1):
            c = ws.cell(row=r, column=col)
            c.font = F_TEXT
            c.border = RAND
            c.alignment = UMBRUCH
            if col <= bis:
                c.fill = FILL_EINGABE


def liste(ws, werte, bereich, warnung=False):
    dv = DataValidation(type="list", formula1='"' + ",".join(werte) + '"', allow_blank=True)
    if warnung:
        dv.errorStyle = "information"
    ws.add_data_validation(dv)
    dv.add(bereich)


def liste_ref(ws, ref, bereich):
    # "information": eigene, neue Codes sind erlaubt (Codebuch ist nur ein Vorschlag)
    dv = DataValidation(type="list", formula1=ref, allow_blank=True, errorStyle="information",
                        error="Neuer Code? Bitte auch im Codebuch ergänzen.", errorTitle="Code nicht im Codebuch")
    ws.add_data_validation(dv)
    dv.add(bereich)


def anleitung(wb):
    ws = wb.active
    ws.title = "Anleitung"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 95
    ws["A1"] = "Erfassung Interviews – Fallstudie 1 SNOCKS"
    ws["A1"].font = F_TITEL
    zeilen = [
        ("Wozu?", "Hier tragen alle Gruppenmitglieder ihre Interviews ein (Leitfaden: fallstudie-1-snocks/interviews/Interviewleitfaden.md). "
                  "Claude wertet die Datei mit auswertung/auswertung.py aus und erzeugt daraus die Grafiken für die Folien."),
        ("Was ausfüllen?", "Nur die HELLGELBEN Zellen. Graue Kursiv-Zeile 2 in jedem Blatt = Beispiel, wird bei der Auswertung ignoriert. "
                           "Blau-graue Spaltenköpfe = Formeln, bitte nicht überschreiben."),
        ("Interview-ID", "Kürzel + Nummer, z. B. HL-01. Dieselbe ID in allen Blättern verwenden. Keine Namen!"),
        ("Blatt Teilnehmende", "1 Zeile pro Interview (Kurzsteckbrief, Teil 1 des Leitfadens)."),
        ("Blatt Grid", "1 Zeile pro Konstrukt (Teil 2). Skala: 1 = linker Pol trifft voll zu … 5 = rechter Pol trifft voll zu. "
                       "Bevorzugter Pol: L oder R. Wichtigkeit: 1, 2 oder 3 bei den Top-3-Konstrukten, sonst leer. "
                       "Kategorie: nachträglich aus dem Codebuch zuordnen (hilft, ähnliche Konstrukte verschiedener Personen zu bündeln)."),
        ("Blatt Laddering", "1 Zeile pro Leiter (Teil 3, drei Leitern pro Interview). Rohnotiz möglichst wörtlich; "
                            "danach jede Stufe einem Code zuordnen (Dropdown = Vorschläge aus dem Codebuch, eigene Codes sind erlaubt). "
                            "Leere Stufen einfach freilassen."),
        ("Blatt Analogien", "1 Zeile pro Interview (Teil 4). Die Begründung („warum?“) ist das Wichtigste."),
        ("Blatt Codebuch", "Vorschläge für Grid-Kategorien und Laddering-Codes. Neue Codes unten ergänzen, damit alle dieselben Begriffe nutzen."),
        ("Datenschutz", "Anonym erfassen, keine Namen oder Kontaktdaten. Einverständnis in Spalte D vermerken."),
        ("Ablage", "Datei nach dem Eintragen ins GitHub-Repo (fallstudie-1-snocks/daten/) bzw. in den Gruppenordner legen."),
    ]
    for i, (a, b) in enumerate(zeilen, start=3):
        ws.cell(row=i, column=1, value=a).font = Font(name=ARIAL, size=10, bold=True)
        c = ws.cell(row=i, column=2, value=b)
        c.font = F_TEXT
        c.alignment = UMBRUCH
        ws.row_dimensions[i].height = 44
    ws.cell(row=14, column=1, value="Legende").font = Font(name=ARIAL, size=10, bold=True)
    for i, (fill, text) in enumerate([(FILL_EINGABE, "Eingabezelle"), (FILL_BEISPIEL, "Beispielzeile (wird ignoriert)"),
                                      (FILL_FORMEL_KOPF, "Formelspalte (nicht überschreiben)")], start=15):
        ws.cell(row=i, column=1).fill = fill
        ws.cell(row=i, column=2, value=text).font = F_TEXT


def teilnehmende(wb):
    ws = wb.create_sheet("Teilnehmende")
    spalten = ["Interview-ID", "Interviewer:in", "Datum", "Einverständnis", "Altersgruppe", "Geschlecht",
               "Wohnort", "Tätigkeit", "SNOCKS bekannt?", "SNOCKS gekauft?", "Wo SNOCKS gekauft?",
               "Wo kaufst du Socken/Unterwäsche sonst?", "Plattform 1", "Plattform 2", "Plattform 3"]
    kopf(ws, spalten, [13, 14, 11, 13, 12, 12, 13, 18, 11, 11, 18, 30, 13, 13, 13])
    beispiel(ws, ["BEISPIEL", "HL", "20.10.2026", "ja", "25-34", "m", "Großstadt", "angestellt", "ja", "ja",
                  "SNOCKS-Shop/App", "dm, H&M, Amazon", "Instagram", "YouTube", "WhatsApp"])
    n = ZEILEN["Teilnehmende"]
    eingabe(ws, len(spalten), n)
    letzte = 2 + n
    liste(ws, ["ja", "nein"], f"D3:D{letzte}")
    liste(ws, ["18-24", "25-34", "35-44", "45-54", "55+"], f"E3:E{letzte}")
    liste(ws, ["w", "m", "d", "k.A."], f"F3:F{letzte}")
    liste(ws, ["Großstadt", "Mittelstadt", "Kleinstadt", "ländlich"], f"G3:G{letzte}")
    liste(ws, ["Studium/Ausbildung", "angestellt", "selbstständig", "sonstiges"], f"H3:H{letzte}")
    liste(ws, ["ja", "nein"], f"I3:J{letzte}")
    liste(ws, ["SNOCKS-Shop/App", "Amazon", "Zalando/About You/Otto", "Laden", "-"], f"K3:K{letzte}", warnung=True)
    liste(ws, ["Instagram", "TikTok", "YouTube", "WhatsApp", "Facebook", "LinkedIn", "Podcasts", "Pinterest",
               "Snapchat", "X/Twitter", "Reddit"], f"M3:O{letzte}", warnung=True)


def grid(wb):
    ws = wb.create_sheet("Grid")
    spalten = (["Interview-ID", "Konstrukt-Nr", "Triade", "Linker Pol (Ähnlichkeit)", "Rechter Pol (Kontrast)",
                "Bevorzugter Pol (L/R)", "Wichtigkeit (Top 3)"] + ELEMENTE +
               ["Kategorie (Codierung)", "SNOCKS orientiert (5 = bevorzugt)", "Ideal orientiert",
                "Abstand Ideal – SNOCKS"])
    kopf(ws, spalten, [13, 10, 8, 30, 30, 11, 11] + [9] * 8 + [24, 13, 11, 12], formel_ab=17)
    ws["Q1"].comment = Comment("Formel: Bewertung so gedreht, dass 5 immer den bevorzugten Pol bedeutet.", "Claude")
    beispiel(ws, ["BEISPIEL", 1, "T1", "hochwertig, hält lange", "billig, schnell kaputt", "L", 1,
                  2, 1, 3, 2, 2, 3, 5, 1, "Qualität & Haltbarkeit"])
    n = ZEILEN["Grid"]
    eingabe(ws, len(spalten), n, bis_spalte=16)
    for r in range(2, 3 + n):
        ws.cell(row=r, column=17, value=f'=IF(OR($A{r}="",H{r}="",$F{r}=""),"",IF($F{r}="L",6-H{r},H{r}))')
        ws.cell(row=r, column=18, value=f'=IF(OR($A{r}="",O{r}="",$F{r}=""),"",IF($F{r}="L",6-O{r},O{r}))')
        ws.cell(row=r, column=19, value=f'=IF(OR(Q{r}="",R{r}=""),"",R{r}-Q{r})')
        for col in (17, 18, 19):
            ws.cell(row=r, column=col).font = F_BEISPIEL if r == 2 else F_TEXT
            ws.cell(row=r, column=col).alignment = Alignment(horizontal="center")
    letzte = 2 + n
    liste(ws, ["T1", "T2", "T3", "T4", "T5"], f"C3:C{letzte}")
    liste(ws, ["L", "R"], f"F3:F{letzte}")
    liste(ws, ["1", "2", "3"], f"G3:G{letzte}")
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True,
                        error="Bitte eine ganze Zahl von 1 bis 5 eintragen.", errorTitle="Skala 1–5")
    ws.add_data_validation(dv)
    dv.add(f"H3:O{letzte}")
    liste_ref(ws, f"=Codebuch!$A$4:$A${3 + len(GRID_KATEGORIEN)}", f"P3:P{letzte}")


def laddering(wb, code_bereiche):
    ws = wb.create_sheet("Laddering")
    spalten = ["Interview-ID", "Leiter-Nr", "Startattribut (wörtlich)", "Rohnotiz (Antwortkette wörtlich)",
               "A – Attribut", "FK – funktionale Konsequenz", "PK – psychosoziale Konsequenz", "W – Wert", "Bemerkung"]
    kopf(ws, spalten, [13, 9, 26, 55, 22, 26, 28, 26, 24])
    beispiel(ws, ["BEISPIEL", 1, "weiche Bio-Baumwolle",
                  "„weil die nicht kratzt und ich weniger schwitze“ → „dann fühl ich mich den ganzen Tag wohl“ → „das ist für mich Lebensqualität“",
                  "Bio-Baumwolle", "angenehmes Tragegefühl", "fühle mich wohl in meiner Haut", "Lebensqualität / Wohlbefinden", None])
    ws.row_dimensions[2].height = 45
    n = ZEILEN["Laddering"]
    eingabe(ws, len(spalten), n)
    letzte = 2 + n
    liste(ws, ["1", "2", "3"], f"B3:B{letzte}")
    for col, ref in zip("EFGH", code_bereiche):
        liste_ref(ws, ref, f"{col}3:{col}{letzte}")


def analogien(wb):
    ws = wb.create_sheet("Analogien")
    spalten = ["Interview-ID", "Mensch (wer/wie?)", "Mensch – warum?", "Prominente Person", "Promi – warum?",
               "Tier", "Tier – warum?", "Auto", "Auto – warum?", "Adjektiv 1", "Adjektiv 2", "Adjektiv 3", "Gefühl"]
    kopf(ws, spalten, [13, 24, 30, 18, 30, 14, 30, 16, 30, 13, 13, 13, 18])
    beispiel(ws, ["BEISPIEL", "Mitte 20, Start-up-Gründer, locker", "duzt alle, macht Witze, wirkt trotzdem zuverlässig",
                  "Joko Winterscheidt", "frech, aber sympathisch", "Golden Retriever", "freundlich und für jeden da",
                  "VW Golf", "praktisch für den Alltag, nichts Besonderes, aber solide", "locker", "bequem", "jung",
                  "Wohlfühlen"])
    ws.row_dimensions[2].height = 45
    eingabe(ws, len(spalten), ZEILEN["Analogien"])


def codebuch(wb):
    ws = wb.create_sheet("Codebuch")
    ws["A1"] = "Codebuch – gemeinsame Begriffe für die Auswertung (Vorschläge, eigene Codes unten ergänzen)"
    ws["A1"].font = F_TITEL
    for col, (titel, breite) in enumerate([("Grid-Kategorie", 26), ("Definition / Beispiele", 50), ("", 3),
                                           ("A – Attribut", 26), ("FK – funktionale Konsequenz", 30),
                                           ("PK – psychosoziale Konsequenz", 34), ("W – Wert", 32)], start=1):
        c = ws.cell(row=3, column=col, value=titel or None)
        if titel:
            c.font = F_KOPF
            c.fill = FILL_KOPF
        ws.column_dimensions[get_column_letter(col)].width = breite
    for i, (kat, defi) in enumerate(GRID_KATEGORIEN, start=4):
        ws.cell(row=i, column=1, value=kat).font = F_TEXT
        ws.cell(row=i, column=2, value=defi).font = F_TEXT
    bereiche = []
    for col, (stufe, codes) in zip([4, 5, 6, 7], CODES.items()):
        for i, code in enumerate(codes, start=4):
            ws.cell(row=i, column=col, value=code).font = F_TEXT
        # Platz für 15 eigene Codes je Stufe
        for i in range(4 + len(codes), 4 + len(codes) + 15):
            ws.cell(row=i, column=col).fill = FILL_EINGABE
        letzte = 3 + len(codes) + 15
        bereiche.append(f"=Codebuch!${get_column_letter(col)}$4:${get_column_letter(col)}${letzte}")
    return bereiche


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    ziel = Path(args[0]) if args else ZIEL
    if ziel.exists() and "--force" not in sys.argv:
        sys.exit(f"{ziel} existiert schon – mit --force überschreiben (Achtung: Daten gehen verloren).")
    wb = Workbook()
    anleitung(wb)
    teilnehmende(wb)
    grid(wb)
    bereiche = codebuch(wb)
    laddering(wb, bereiche)
    analogien(wb)
    wb.move_sheet("Codebuch", offset=10)
    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = 110
    ziel.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ziel)
    print(f"geschrieben: {ziel}")


if __name__ == "__main__":
    main()
